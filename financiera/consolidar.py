# -*- coding: utf-8 -*-
"""
Consolidación de la cartera de créditos de la financiera Yamaha (Pitalito/Garzón).

Lee el archivo original FINANCIERA_2025.xlsx y genera:
  - FINANCIERA_2025_CONSOLIDADO.xlsx  (hojas: RESUMEN, CLIENTES, CREDITOS, CUOTAS, MORA, HALLAZGOS)
  - csv/clientes.csv, csv/creditos.csv, csv/cuotas.csv  (UTF-8 BOM, separador ';')

Fuente autoritativa: hoja LETRAS (una fila por cuota de cada crédito).
Las 31 hojas de caja diaria se usan para extraer los DESEMBOLSOS (capital prestado)
y para validar el cuadre contra lo registrado en LETRAS.

Toda corrección sobre los datos originales está declarada explícitamente en
DATE_FIXES / NAME_FIXES y queda documentada en la hoja HALLAZGOS. Nada se
modifica en silencio.

Uso:
    python3 consolidar.py [--origen RUTA_XLSX] [--corte AAAA-MM-DD] [--salida CARPETA]
"""
import argparse
import csv
import datetime
import os
import re
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ORIGEN_DEFAULT = os.path.join(BASE_DIR, 'origen', 'FINANCIERA_2025.xlsx')
CORTE_DEFAULT = '2026-07-15'  # fecha de corte para estados de mora

HOJAS_NO_DIARIAS = {'LETRAS', 'NORMA', 'SIMULADOR', 'APORTES', 'Data'}

# Correcciones documentadas (ver hoja HALLAZGOS del archivo generado).
# Fechas: errores de digitación del año, corregidos manteniendo mes y día.
DATE_FIXES = [
    # (hoja, fila, campo, valor_original, valor_corregido, justificación)
    ('LETRAS', 13, 'FECHA_PAGO', datetime.date(2005, 3, 28), datetime.date(2025, 3, 28),
     'Año 2005 imposible: el crédito #2234 inició en 2025; cuotas vecinas pagadas en 2025.'),
    ('LETRAS', 856, 'FECHA_VENCIMIENTO', datetime.date(2006, 6, 30), datetime.date(2026, 6, 30),
     'Año 2006 imposible: cuota 10 del crédito #2674; la cuota 9 vence 2026-05-30 y la 11 2026-07-30.'),
]

# Nombres: unificación de grafías de la misma persona entre LETRAS y hojas de caja.
# Se toma como forma canónica el apellido real (VALENCIA, PERAFAN).
NAME_FIXES = {
    'NOHORA ELCY MADRIGAL VALENCUA': 'NOHORA ELCY MADRIGAL VALENCIA',
    'ANYI PAOLA PEFARAN HOYOS': 'ANYI PAOLA PERAFAN HOYOS',
}

HALLAZGOS_ADICIONALES = [
    ('CRUCE DE NOMBRES EN CAJA',
     'Hoja 7-MARZO',
     'Un ingreso único de $5,472,500 (cuotas 3 a 12 del crédito #2205 de ANDRES HENRY HERNANDEZ DELGADO) '
     'quedó anotado en caja como "PAGO LETRAS JUAN CARLOS CUELLAR ALAPE". En LETRAS está aplicado '
     'correctamente al #2205. Los pagos mensuales de $672,300 de JUAN CARLOS CUELLAR corresponden a su '
     'propio crédito #2154. No requiere corrección: LETRAS es la fuente autoritativa y el total cuadra al peso.'),
    ('CRUCE DE NOMBRES EN CAJA',
     'Hoja 14-JULIO',
     'Abono de $3,900 registrado en caja a nombre de EDI ALEXANDER JIMENEZ VARGAS pero aplicado en LETRAS a '
     'EDWIN ARNOLD ORTIZ PINO (#2948, cuota 7, abono parcial). Verificar con los recibos físicos. '
     'Impacto: $3,900.'),
    ('ABONOS PARCIALES VIVOS',
     'LETRAS (créditos #78, #2640, #2731, #2948)',
     '4 cuotas de créditos activos tienen un ABONO PARCIAL registrado (la caja los rotula "ABONO LETRA"): '
     '#78 JUAN CARLOS NUÑEZ PARRA cuota 3: abonó $170,600 de $614,700 (debe $444,100); '
     '#2640 JULIAN FERNANDO ACHURY BRAVO cuota 7: abonó $100,304 de $789,900 (debe $689,596); '
     '#2731 MAURICIO TORRES SERRATO cuota 6: abonó $300,000 de $1,096,900 (debe $796,900); '
     '#2948 EDWIN ARNOLD ORTIZ PINO cuota 7: abonó $3,900 de $298,000 (debe $294,100). '
     'Estas cuotas figuran con estado ABONO PARCIAL y su faltante ($2,224,696 en total) está INCLUIDO en el '
     'saldo por cobrar. La convención del libro es completar el valor en la misma celda cuando el cliente '
     'termina de pagar la cuota.'),
    ('HOMÓNIMOS - NO FUSIONAR',
     'LETRAS',
     'CARLA FERNANDA MUÑOZ ROJAS (crédito #3007, 2026) y MARIA FERNANDA MUÑOZ ROJAS (crédito #2246, 2025, '
     'cancelado) son personas DISTINTAS con créditos distintos. No unificar.'),
    ('CLIENTE REPITENTE',
     'LETRAS',
     'JESUS MARIA CASTAÑEDA IJAJI tiene 2 créditos (#2286 y #2648), ambos cancelados. Es el único cliente '
     'con más de un crédito.'),
    ('MORA CRÍTICA',
     'LETRAS',
     'CATALINA CEDEÑO PLAZAS (#2473): sin pagos desde 2025-09-08. Revisar gestión de cobro.'),
    ('CANCELACIÓN ANTICIPADA ATÍPICA',
     'LETRAS crédito #2205',
     'ANDRES HENRY HERNANDEZ DELGADO (#2205): el crédito completo se liquidó en un solo día (2025-03-05, '
     '19 días ANTES de vencer la primera cuota): cuota 1 a valor pleno, cuota 2 por $644,766 y cuotas 3-12 '
     'con descuento ($547,250 vs letra de $679,600). Confirmado contra caja (hoja 7-MARZO). Figura con '
     'CIERRE=ANTICIPADO.'),
    ('FUERA DEL ALCANCE - NORMA',
     'Hoja NORMA',
     'Préstamo aparte de 18 cuotas de $1,000,000 (2 pagadas por NEQUI: 2026-04-06 y 2026-05-09; '
     'saldo $16,000,000). No identifica número de crédito ni ciudad. NO se incluyó en CLIENTES/CREDITOS; '
     'decidir si se registra en el sistema.'),
    ('FUERA DEL ALCANCE - APORTES',
     'Hoja APORTES',
     'Movimientos de capital con MOTO JAPONESA (CxP/abonos). Último saldo en caja registrado: '
     '$179,886,118 al 2026-07-14. Es contabilidad interna, no cartera de clientes.'),
]

# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def norm_name(s):
    """Normaliza espacios y mayúsculas de un nombre."""
    return re.sub(r'\s+', ' ', str(s).strip().upper())


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


# ---------------------------------------------------------------------------
# Lectura de LETRAS
# ---------------------------------------------------------------------------

def leer_letras(wb):
    """Lee la hoja LETRAS. Devuelve la lista de cuotas con correcciones aplicadas."""
    ws = wb['LETRAS']
    fixes = {(f[1], f[2]): f for f in DATE_FIXES if f[0] == 'LETRAS'}
    cuotas = []
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        due, num, city, name, valor, pagado, fpago, ncuota = (list(r) + [None] * 8)[:8]
        if num is None and name is None:
            continue
        due, fpago = as_date(due), as_date(fpago)
        # correcciones de fecha documentadas
        fx = fixes.get((i, 'FECHA_VENCIMIENTO'))
        if fx:
            assert due == fx[3], f'Fila {i}: se esperaba {fx[3]} y hay {due}; revisar DATE_FIXES'
            due = fx[4]
        fx = fixes.get((i, 'FECHA_PAGO'))
        if fx:
            assert fpago == fx[3], f'Fila {i}: se esperaba {fx[3]} y hay {fpago}; revisar DATE_FIXES'
            fpago = fx[4]
        name = norm_name(name)
        name = NAME_FIXES.get(name, name)
        cuotas.append({
            'fila_origen': i,
            'credito': num,
            'ciudad': norm_name(city),
            'cliente': name,
            'vence': due,
            'valor': float(valor) if valor is not None else 0.0,
            'pagado': float(pagado) if pagado not in (None, '', 0) else None,
            'fecha_pago': fpago,
            'cuota': int(ncuota),
        })
    return cuotas


# ---------------------------------------------------------------------------
# Lectura de hojas de caja diaria (desembolsos + validación de ingresos)
# ---------------------------------------------------------------------------

def leer_caja(wb):
    """Extrae de cada hoja de caja los ingresos por letras y los desembolsos."""
    ingresos, desembolsos = [], []
    hojas = [n for n in wb.sheetnames if n not in HOJAS_NO_DIARIAS]
    for sn in hojas:
        ws = wb[sn]
        seccion = None  # None | 'ingresos' | 'egresos'
        for r in ws.iter_rows(values_only=True):
            c0 = norm_name(r[0]) if r[0] is not None else ''
            c2 = norm_name(r[2]) if len(r) > 2 and r[2] is not None else ''
            c4 = norm_name(r[4]) if len(r) > 4 and r[4] is not None else ''
            if c0 == 'LETRAS' and 'CUOTA' in c2:
                seccion = 'ingresos'
                continue
            if 'TOTAL INGRESOS' in c2 or 'TOTAL INGRESOS' in c0:
                seccion = None
                continue
            if 'PAGOS EFECTUADOS' in c0:
                seccion = 'egresos'
                continue
            if 'TOTAL PAGOS' in c0:
                seccion = None
                continue
            monto = r[11] if len(r) > 11 else None
            if not isinstance(monto, (int, float)) or monto == 0 or not c4:
                continue
            if seccion == 'ingresos':
                ingresos.append({'hoja': sn, 'concepto': c4, 'total': float(monto)})
            elif seccion == 'egresos' and c4.startswith('DESEMBOLSO'):
                nombre = NAME_FIXES.get(c4.replace('DESEMBOLSO', '').strip(),
                                        c4.replace('DESEMBOLSO', '').strip())
                desembolsos.append({
                    'hoja': sn,
                    'fecha': as_date(r[0]),
                    'cliente': nombre,
                    'monto': float(monto),
                })
    return ingresos, desembolsos


def clasificar_ingreso(concepto):
    """Clasifica un concepto de la sección de ingresos de caja."""
    if re.match(r'^(PAGO|PGO)\s+LE[TY]?T?RA', concepto):
        return 'PAGO_LETRA'
    if concepto.startswith('ABONO LETRA'):
        return 'ABONO_LETRA'
    if concepto.startswith('SALDO LETRA') or concepto.startswith('SALDO DEUDA'):
        return 'SALDO_LETRA'
    if 'DATACREDITO' in concepto:
        return 'DATACREDITO'
    if 'CXP' in concepto or 'MOTO JAPONESA' in concepto:
        return 'APORTE_CAPITAL'
    return 'OTRO'


def total_letras_en_caja(ingresos, nombres_clientes):
    """Suma todo ingreso de caja atribuible a pagos de letras de clientes.

    Incluye PAGO/ABONO/SALDO LETRA y filas cuyo concepto es solo el nombre
    del cliente (casos documentados: pagos registrados sin prefijo).
    """
    total = 0.0
    detalle_otros = []
    for p in ingresos:
        k = clasificar_ingreso(p['concepto'])
        if k in ('PAGO_LETRA', 'ABONO_LETRA', 'SALDO_LETRA'):
            total += p['total']
        elif k == 'OTRO':
            nombre = NAME_FIXES.get(p['concepto'], p['concepto'])
            if nombre in nombres_clientes:
                total += p['total']
                detalle_otros.append(p)
    return total, detalle_otros


# ---------------------------------------------------------------------------
# Consolidación
# ---------------------------------------------------------------------------

def consolidar(cuotas, desembolsos, corte):
    """Agrupa cuotas por crédito y clientes; asigna desembolsos y estados."""
    por_credito = defaultdict(list)
    for c in cuotas:
        por_credito[(c['credito'], c['cliente'])].append(c)

    creditos = []
    for (num, cliente), rows in por_credito.items():
        rows.sort(key=lambda x: x['cuota'])
        # integridad de la secuencia de cuotas
        secuencia = [r['cuota'] for r in rows]
        assert secuencia == list(range(1, len(rows) + 1)), \
            f'Crédito #{num} {cliente}: secuencia de cuotas rota: {secuencia}'
        valores = {r['valor'] for r in rows}
        assert len(valores) == 1, f'Crédito #{num} {cliente}: valor de letra inconsistente {valores}'

        con_pago = [r for r in rows if r['pagado'] is not None]
        sin_pago = [r for r in rows if r['pagado'] is None]
        abierto = bool(sin_pago)
        # En un crédito ABIERTO, una cuota con pago menor a la letra es un abono
        # parcial: el faltante sigue siendo deuda (la caja los rotula ABONO LETRA).
        # En un crédito ya liquidado, pagar menos que la letra es el descuento del
        # cierre anticipado y no deja deuda.
        parciales = [r for r in con_pago if abierto and r['pagado'] < r['valor']]
        pagadas = [r for r in con_pago if not (abierto and r['pagado'] < r['valor'])]
        deudoras = sin_pago + parciales
        vencidas = [r for r in deudoras if r['vence'] and r['vence'] < corte]
        total_pagado = sum(r['pagado'] for r in con_pago)
        saldo = sum(r['valor'] for r in sin_pago) + sum(r['valor'] - r['pagado'] for r in parciales)

        # anotar cada cuota para las hojas CUOTAS / cuotas.csv
        for r in rows:
            if r['pagado'] is None:
                r['saldo_cuota'] = r['valor']
                r['estado_cuota'] = 'VENCIDA' if r['vence'] and r['vence'] < corte else 'PENDIENTE'
            elif abierto and r['pagado'] < r['valor']:
                r['saldo_cuota'] = r['valor'] - r['pagado']
                r['estado_cuota'] = 'ABONO PARCIAL'
            else:
                r['saldo_cuota'] = 0.0
                r['estado_cuota'] = 'PAGADA'

        if not deudoras:
            estado = 'CANCELADO'
        elif vencidas:
            estado = 'EN MORA'
        else:
            estado = 'AL DIA'

        # Cierre ANTICIPADO: el último día de pago liquidó de una vez un tramo
        # final de >=2 cuotas que aún no habían vencido o que se pagaron con
        # descuento (pago del saldo). Pagar cada mes su cuota (incluso unos días
        # antes) deja tramos de 1 sola cuota y queda como NORMAL.
        cierre = ''
        if estado == 'CANCELADO':
            cierre = 'NORMAL'
            fecha_ult = max(r['fecha_pago'] for r in rows if r['fecha_pago'])
            cola = 0
            for r in reversed(rows):
                if r['fecha_pago'] == fecha_ult and (r['pagado'] < r['valor'] or
                                                     (r['vence'] and r['vence'] > fecha_ult)):
                    cola += 1
                else:
                    break
            if cola >= 2:
                cierre = 'ANTICIPADO'

        creditos.append({
            'credito': num,
            'cliente': cliente,
            'ciudad': rows[0]['ciudad'],
            'valor_letra': rows[0]['valor'],
            'plazo': len(rows),
            'primera_cuota_vence': min(r['vence'] for r in rows),
            'ultima_cuota_vence': max(r['vence'] for r in rows),
            'cuotas_pagadas': len(pagadas),
            'cuotas_abono_parcial': len(parciales),
            'cuotas_pendientes': len(sin_pago),
            'cuotas_vencidas': len(vencidas),
            'total_pagado': total_pagado,
            'saldo_pendiente': saldo,
            'proxima_cuota_vence': min((r['vence'] for r in deudoras), default=None),
            'ultimo_pago': max((r['fecha_pago'] for r in con_pago if r['fecha_pago']), default=None),
            'dias_mora': (corte - min(r['vence'] for r in vencidas)).days if vencidas else 0,
            'estado': estado,
            'cierre': cierre,
            'fecha_desembolso': None,
            'monto_desembolso': None,
        })

    # asignación de desembolsos: por cliente, en orden cronológico
    des_por_cliente = defaultdict(list)
    for d in desembolsos:
        des_por_cliente[d['cliente']].append(d)
    for lst in des_por_cliente.values():
        lst.sort(key=lambda d: d['fecha'] or datetime.date.min)

    creditos_por_cliente = defaultdict(list)
    for c in creditos:
        creditos_por_cliente[c['cliente']].append(c)

    sin_desembolso = []
    for cliente, lst in creditos_por_cliente.items():
        lst.sort(key=lambda c: c['primera_cuota_vence'])
        dl = des_por_cliente.get(cliente, [])
        for i, c in enumerate(lst):
            if i < len(dl):
                c['fecha_desembolso'] = dl[i]['fecha']
                c['monto_desembolso'] = dl[i]['monto']
            else:
                sin_desembolso.append(c)

    desembolsos_sueltos = []
    for cliente, dl in des_por_cliente.items():
        n = len(creditos_por_cliente.get(cliente, []))
        desembolsos_sueltos.extend(dl[n:])

    # clientes
    clientes = []
    # Orden estable: fecha de primera cuota y, en empate, nombre. Así los IDs
    # no dependen del orden físico de las filas en la hoja LETRAS. Aun así los
    # C-ID son de ESTE corte: la clave definitiva del cliente será la cédula.
    orden = sorted(creditos_por_cliente.items(),
                   key=lambda kv: (min(c['primera_cuota_vence'] for c in kv[1]), kv[0]))
    for idx, (nombre, lst) in enumerate(orden, 1):
        estados = {c['estado'] for c in lst}
        if 'EN MORA' in estados:
            estado = 'EN MORA'
        elif 'AL DIA' in estados:
            estado = 'AL DIA'
        else:
            estado = 'CANCELADO'
        clientes.append({
            'id': f'C{idx:03d}',
            'nombre': nombre,
            'ciudad': lst[0]['ciudad'],
            'cedula': '',    # pendiente: documentos físicos
            'telefono': '',  # pendiente: documentos físicos
            'n_creditos': len(lst),
            'creditos': ', '.join(str(c['credito']) for c in lst),
            'total_desembolsado': sum(c['monto_desembolso'] or 0 for c in lst),
            'total_pagado': sum(c['total_pagado'] for c in lst),
            'saldo_pendiente': sum(c['saldo_pendiente'] for c in lst),
            'estado': estado,
        })
    id_cliente = {c['nombre']: c['id'] for c in clientes}
    for c in creditos:
        c['id_cliente'] = id_cliente[c['cliente']]

    creditos.sort(key=lambda c: (c['primera_cuota_vence'], str(c['credito'])))
    return clientes, creditos, sin_desembolso, desembolsos_sueltos


# ---------------------------------------------------------------------------
# Salidas
# ---------------------------------------------------------------------------
HDR_FILL = PatternFill('solid', fgColor='1F4E79')
HDR_FONT = Font(bold=True, color='FFFFFF')
MONEY = '#,##0'


def _hoja(wb, titulo, encabezados, filas, money_cols=(), date_cols=(), widths=None):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for fila in filas:
        ws.append(fila)
    for col in money_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = MONEY
    for col in date_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = 'YYYY-MM-DD'
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def escribir_excel(ruta, corte, clientes, creditos, cuotas, resumen, hallazgos):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # RESUMEN
    ws = wb.create_sheet('RESUMEN')
    ws.column_dimensions['A'].width = 52
    ws.column_dimensions['B'].width = 22
    ws.append(['CONSOLIDADO CARTERA FINANCIERA YAMAHA', ''])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append(['Fecha de corte', corte])
    ws.append([])
    for k, v in resumen:
        ws.append([k, v])
        if isinstance(v, (int, float)) and not isinstance(v, bool) and abs(v) > 1000:
            ws.cell(ws.max_row, 2).number_format = MONEY
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        if str(row[0].value or '').startswith('---'):
            row[0].font = Font(bold=True)

    # CLIENTES
    _hoja(wb, 'CLIENTES',
          ['ID', 'NOMBRE', 'CIUDAD', 'CEDULA', 'TELEFONO', 'N_CREDITOS', 'CREDITOS',
           'TOTAL_DESEMBOLSADO', 'TOTAL_PAGADO', 'SALDO_PENDIENTE', 'ESTADO'],
          [[c['id'], c['nombre'], c['ciudad'], c['cedula'], c['telefono'], c['n_creditos'],
            c['creditos'], c['total_desembolsado'], c['total_pagado'], c['saldo_pendiente'],
            c['estado']] for c in clientes],
          money_cols=(8, 9, 10),
          widths=[6, 40, 12, 12, 12, 10, 12, 17, 15, 15, 12])

    # CREDITOS
    _hoja(wb, 'CREDITOS',
          ['CREDITO', 'ID_CLIENTE', 'CLIENTE', 'CIUDAD', 'FECHA_DESEMBOLSO', 'MONTO_DESEMBOLSO',
           'VALOR_LETRA', 'PLAZO_CUOTAS', 'PRIMERA_CUOTA', 'ULTIMA_CUOTA', 'CUOTAS_PAGADAS',
           'CUOTAS_ABONO_PARCIAL', 'CUOTAS_PENDIENTES', 'CUOTAS_VENCIDAS', 'TOTAL_PAGADO',
           'SALDO_PENDIENTE', 'PROXIMA_CUOTA', 'ULTIMO_PAGO', 'DIAS_MORA', 'ESTADO', 'CIERRE'],
          [[c['credito'], c['id_cliente'], c['cliente'], c['ciudad'], c['fecha_desembolso'],
            c['monto_desembolso'], c['valor_letra'], c['plazo'], c['primera_cuota_vence'],
            c['ultima_cuota_vence'], c['cuotas_pagadas'], c['cuotas_abono_parcial'],
            c['cuotas_pendientes'], c['cuotas_vencidas'], c['total_pagado'], c['saldo_pendiente'],
            c['proxima_cuota_vence'], c['ultimo_pago'], c['dias_mora'], c['estado'],
            c['cierre']] for c in creditos],
          money_cols=(6, 7, 15, 16),
          date_cols=(5, 9, 10, 17, 18),
          widths=[9, 10, 40, 11, 16, 16, 12, 9, 12, 12, 9, 10, 9, 9, 14, 14, 12, 12, 10, 11, 12])

    # CUOTAS
    id_por_nombre = {c['nombre']: c['id'] for c in clientes}
    filas_cuotas = []
    for q in sorted(cuotas, key=lambda x: (str(x['credito']), x['cuota'])):
        atraso = None
        if q['estado_cuota'] == 'PAGADA' and q['fecha_pago'] and q['vence']:
            atraso = max(0, (q['fecha_pago'] - q['vence']).days)
        filas_cuotas.append([q['credito'], id_por_nombre[q['cliente']], q['cliente'], q['cuota'],
                             q['vence'], q['valor'], q['pagado'], q['fecha_pago'],
                             q['saldo_cuota'], atraso, q['estado_cuota'], q['fila_origen']])
    _hoja(wb, 'CUOTAS',
          ['CREDITO', 'ID_CLIENTE', 'CLIENTE', 'CUOTA', 'FECHA_VENCIMIENTO', 'VALOR_LETRA',
           'VALOR_PAGADO', 'FECHA_PAGO', 'SALDO_CUOTA', 'DIAS_ATRASO_PAGO', 'ESTADO',
           'FILA_ORIGEN_LETRAS'],
          filas_cuotas,
          money_cols=(6, 7, 9),
          date_cols=(5, 8),
          widths=[9, 10, 40, 7, 16, 13, 13, 12, 13, 14, 14, 14])

    # MORA
    en_mora = sorted([c for c in creditos if c['estado'] == 'EN MORA'],
                     key=lambda c: -c['dias_mora'])
    _hoja(wb, 'MORA',
          ['CREDITO', 'CLIENTE', 'CIUDAD', 'CUOTAS_VENCIDAS', 'DIAS_MORA', 'VALOR_LETRA',
           'SALDO_PENDIENTE', 'ULTIMO_PAGO', 'CUOTA_VENCIDA_MAS_ANTIGUA'],
          [[c['credito'], c['cliente'], c['ciudad'], c['cuotas_vencidas'], c['dias_mora'],
            c['valor_letra'], c['saldo_pendiente'], c['ultimo_pago'], c['proxima_cuota_vence']]
           for c in en_mora],
          money_cols=(6, 7),
          date_cols=(8, 9),
          widths=[9, 40, 12, 15, 10, 13, 15, 12, 22])

    # HALLAZGOS
    filas_h = []
    for hoja, fila, campo, orig, corr, just in DATE_FIXES:
        filas_h.append(['CORRECCIÓN DE FECHA', f'{hoja} fila {fila} ({campo})',
                        str(orig), str(corr), just])
    for orig, corr in NAME_FIXES.items():
        filas_h.append(['UNIFICACIÓN DE NOMBRE', 'LETRAS / hojas de caja', orig, corr,
                        'Misma persona escrita distinto entre hojas; se usa el apellido real.'])
    for tipo, ubic, desc in HALLAZGOS_ADICIONALES:
        filas_h.append([tipo, ubic, '', '', desc])
    _hoja(wb, 'HALLAZGOS',
          ['TIPO', 'UBICACIÓN', 'VALOR ORIGINAL', 'VALOR USADO', 'DETALLE / JUSTIFICACIÓN'],
          filas_h,
          widths=[24, 26, 26, 26, 110])
    for row in wb['HALLAZGOS'].iter_rows(min_row=2):
        row[4].alignment = Alignment(wrap_text=True, vertical='top')

    wb.save(ruta)


def escribir_csvs(carpeta, clientes, creditos, cuotas):
    os.makedirs(carpeta, exist_ok=True)

    def dump(nombre, encabezados, filas):
        with open(os.path.join(carpeta, nombre), 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(encabezados)
            w.writerows(filas)

    dump('clientes.csv',
         ['id', 'nombre', 'ciudad', 'cedula', 'telefono', 'n_creditos', 'creditos',
          'total_desembolsado', 'total_pagado', 'saldo_pendiente', 'estado'],
         [[c['id'], c['nombre'], c['ciudad'], c['cedula'], c['telefono'], c['n_creditos'],
           c['creditos'], int(c['total_desembolsado']), int(c['total_pagado']),
           int(c['saldo_pendiente']), c['estado']] for c in clientes])

    dump('creditos.csv',
         ['credito', 'id_cliente', 'cliente', 'ciudad', 'fecha_desembolso', 'monto_desembolso',
          'valor_letra', 'plazo_cuotas', 'primera_cuota', 'ultima_cuota', 'cuotas_pagadas',
          'cuotas_abono_parcial', 'cuotas_pendientes', 'cuotas_vencidas', 'total_pagado',
          'saldo_pendiente', 'proxima_cuota', 'ultimo_pago', 'dias_mora', 'estado', 'cierre'],
         [[c['credito'], c['id_cliente'], c['cliente'], c['ciudad'],
           c['fecha_desembolso'] or '', int(c['monto_desembolso'] or 0), int(c['valor_letra']),
           c['plazo'], c['primera_cuota_vence'], c['ultima_cuota_vence'], c['cuotas_pagadas'],
           c['cuotas_abono_parcial'], c['cuotas_pendientes'], c['cuotas_vencidas'],
           int(c['total_pagado']), int(c['saldo_pendiente']), c['proxima_cuota_vence'] or '',
           c['ultimo_pago'] or '', c['dias_mora'], c['estado'], c['cierre']] for c in creditos])

    id_por_nombre = {c['nombre']: c['id'] for c in clientes}
    dump('cuotas.csv',
         ['credito', 'id_cliente', 'cliente', 'cuota', 'fecha_vencimiento', 'valor_letra',
          'valor_pagado', 'fecha_pago', 'saldo_cuota', 'estado', 'fila_origen_letras'],
         [[q['credito'], id_por_nombre[q['cliente']], q['cliente'], q['cuota'], q['vence'],
           int(q['valor']), int(q['pagado']) if q['pagado'] is not None else '',
           q['fecha_pago'] or '', int(q['saldo_cuota']), q['estado_cuota'], q['fila_origen']]
          for q in sorted(cuotas, key=lambda x: (str(x['credito']), x['cuota']))])


# ---------------------------------------------------------------------------
# Principal
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--origen', default=ORIGEN_DEFAULT)
    ap.add_argument('--corte', default=CORTE_DEFAULT)
    ap.add_argument('--salida', default=BASE_DIR)
    args = ap.parse_args()
    corte = datetime.date.fromisoformat(args.corte)

    wb = openpyxl.load_workbook(args.origen, data_only=True)
    cuotas = leer_letras(wb)
    ingresos, desembolsos = leer_caja(wb)
    clientes, creditos, sin_des, des_sueltos = consolidar(cuotas, desembolsos, corte)

    # validaciones internas duras
    assert not sin_des, f'Créditos sin desembolso: {[(c["credito"], c["cliente"]) for c in sin_des]}'
    assert not des_sueltos, f'Desembolsos sin crédito: {des_sueltos}'

    total_pagado = sum(c['total_pagado'] for c in creditos)
    total_caja, otros_caja = total_letras_en_caja(ingresos, {c['cliente'] for c in cuotas})
    diferencia_caja = total_pagado - total_caja

    n_mora = sum(1 for c in creditos if c['estado'] == 'EN MORA')
    n_aldia = sum(1 for c in creditos if c['estado'] == 'AL DIA')
    n_canc = sum(1 for c in creditos if c['estado'] == 'CANCELADO')

    resumen = [
        ('--- CARTERA ---', ''),
        ('Clientes', len(clientes)),
        ('Créditos', len(creditos)),
        ('Cuotas registradas', len(cuotas)),
        ('Créditos cancelados', n_canc),
        ('Créditos al día', n_aldia),
        ('Créditos en mora', n_mora),
        ('--- DINERO ---', ''),
        ('Total desembolsado (capital prestado)', sum(c['monto_desembolso'] or 0 for c in creditos)),
        ('Total recaudado (pagos y abonos)', total_pagado),
        ('Saldo por cobrar (pendientes + faltante de abonos parciales)',
         sum(c['saldo_pendiente'] for c in creditos)),
        ('   del cual: faltante de 4 cuotas con abono parcial',
         sum(c['saldo_pendiente'] for c in creditos)
         - sum(q['valor'] for q in cuotas if q['estado_cuota'] in ('PENDIENTE', 'VENCIDA'))),
        ('Saldo por cobrar solo créditos EN MORA',
         sum(c['saldo_pendiente'] for c in creditos if c['estado'] == 'EN MORA')),
        ('--- PRUEBA DE CUADRE ---', ''),
        ('Pagos de letras según hojas de caja diaria', total_caja),
        ('Pagos de letras según libro LETRAS', total_pagado),
        ('Diferencia (debe ser 0)', diferencia_caja),
        ('--- NOTAS ---', ''),
        ('Cédulas y teléfonos', 'Pendientes de los documentos físicos'),
        ('Fuente autoritativa', 'Hoja LETRAS del archivo original'),
        ('Generado por', 'financiera/consolidar.py'),
    ]

    os.makedirs(args.salida, exist_ok=True)
    ruta_xlsx = os.path.join(args.salida, 'FINANCIERA_2025_CONSOLIDADO.xlsx')
    escribir_excel(ruta_xlsx, args.corte, clientes, creditos, cuotas, resumen, HALLAZGOS_ADICIONALES)
    escribir_csvs(os.path.join(args.salida, 'csv'), clientes, creditos, cuotas)

    print(f'Clientes: {len(clientes)} | Créditos: {len(creditos)} | Cuotas: {len(cuotas)}')
    print(f'Estados: CANCELADO={n_canc} AL DIA={n_aldia} EN MORA={n_mora}')
    print(f'Total desembolsado: ${sum(c["monto_desembolso"] or 0 for c in creditos):,.0f}')
    print(f'Total recaudado:    ${total_pagado:,.0f}')
    print(f'Saldo por cobrar:   ${sum(c["saldo_pendiente"] for c in creditos):,.0f}')
    print(f'Cuadre caja vs LETRAS: ${diferencia_caja:,.0f} (debe ser 0)')
    if otros_caja:
        print('Ingresos sin prefijo estándar contados como pago de letra:')
        for p in otros_caja:
            print(f'  {p["hoja"]:16} {p["concepto"][:60]:60} ${p["total"]:,.0f}')
    print(f'\nGenerado: {ruta_xlsx}')
    print(f'Generado: {os.path.join(args.salida, "csv")}/clientes.csv, creditos.csv, cuotas.csv')


if __name__ == '__main__':
    main()
