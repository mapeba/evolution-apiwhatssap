# -*- coding: utf-8 -*-
"""
Verificación independiente de los archivos generados por consolidar.py.

NO reutiliza variables en memoria: relee desde disco el Excel original,
el Excel consolidado y los CSVs, y valida todos los invariantes. Si algo
no cuadra, termina con error y lo reporta.

Uso:
    python3 verificar.py [--origen RUTA_XLSX] [--corte AAAA-MM-DD] [--salida CARPETA]
"""
import argparse
import csv
import datetime
import os
import re
import sys
from collections import defaultdict

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Valores esperados, confirmados en el análisis del archivo original
# (véase README.md). Si el archivo de origen cambia, actualizarlos.
ESPERADO = {
    'clientes': 132,
    'creditos': 133,
    'cuotas': 1601,
    'total_recaudado': 662_178_582,
    # incluye el faltante de las 4 cuotas con abono parcial ($2,224,696):
    # #78 ($444,100), #2640 ($689,596), #2731 ($796,900), #2948 ($294,100)
    'saldo_por_cobrar': 433_414_843,
    'abonos_parciales': 4,
    'total_desembolsado': 883_710_100,
    'cancelados': 47,
    'al_dia': 61,
    'en_mora': 25,
}

errores = []


def check(cond, msg):
    estado = 'OK ' if cond else 'FALLO'
    print(f'  [{estado}] {msg}')
    if not cond:
        errores.append(msg)


def leer_csv(ruta):
    with open(ruta, encoding='utf-8-sig') as f:
        return list(csv.DictReader(f, delimiter=';'))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--origen', default=os.path.join(BASE_DIR, 'origen', 'FINANCIERA_2025.xlsx'))
    ap.add_argument('--corte', default='2026-07-15')
    ap.add_argument('--salida', default=BASE_DIR)
    args = ap.parse_args()
    corte = datetime.date.fromisoformat(args.corte)

    ruta_xlsx = os.path.join(args.salida, 'FINANCIERA_2025_CONSOLIDADO.xlsx')
    clientes = leer_csv(os.path.join(args.salida, 'csv', 'clientes.csv'))
    creditos = leer_csv(os.path.join(args.salida, 'csv', 'creditos.csv'))
    cuotas = leer_csv(os.path.join(args.salida, 'csv', 'cuotas.csv'))

    print('1. Conteos de los CSVs generados')
    check(len(clientes) == ESPERADO['clientes'], f"clientes.csv: {len(clientes)} == {ESPERADO['clientes']}")
    check(len(creditos) == ESPERADO['creditos'], f"creditos.csv: {len(creditos)} == {ESPERADO['creditos']}")
    check(len(cuotas) == ESPERADO['cuotas'], f"cuotas.csv: {len(cuotas)} == {ESPERADO['cuotas']}")

    print('2. Totales de dinero (CSVs)')
    tot_pag = sum(int(q['valor_pagado']) for q in cuotas if q['valor_pagado'])
    tot_pend = sum(int(q['saldo_cuota']) for q in cuotas)
    tot_des = sum(int(c['monto_desembolso']) for c in creditos)
    n_abonos = sum(1 for q in cuotas if q['estado'] == 'ABONO PARCIAL')
    check(tot_pag == ESPERADO['total_recaudado'],
          f"Σ valor_pagado en cuotas.csv = ${tot_pag:,} == ${ESPERADO['total_recaudado']:,}")
    check(tot_pend == ESPERADO['saldo_por_cobrar'],
          f"Σ saldo_cuota en cuotas.csv = ${tot_pend:,} == ${ESPERADO['saldo_por_cobrar']:,}")
    check(tot_des == ESPERADO['total_desembolsado'],
          f"Σ desembolsos en creditos.csv = ${tot_des:,} == ${ESPERADO['total_desembolsado']:,}")
    check(n_abonos == ESPERADO['abonos_parciales'],
          f"cuotas con ABONO PARCIAL: {n_abonos} == {ESPERADO['abonos_parciales']}")

    print('3. Coherencia interna creditos.csv vs cuotas.csv')
    por_credito = defaultdict(list)
    for q in cuotas:
        por_credito[q['credito']].append(q)
    malos = 0
    for c in creditos:
        qs = por_credito[c['credito']]
        pagadas = [q for q in qs if q['estado'] == 'PAGADA']
        parciales = [q for q in qs if q['estado'] == 'ABONO PARCIAL']
        pendientes = [q for q in qs if q['estado'] in ('PENDIENTE', 'VENCIDA')]
        okc = (len(qs) == int(c['plazo_cuotas'])
               and len(pagadas) == int(c['cuotas_pagadas'])
               and len(parciales) == int(c['cuotas_abono_parcial'])
               and len(pendientes) == int(c['cuotas_pendientes'])
               and len(pagadas) + len(parciales) + len(pendientes) == len(qs)
               and sum(int(q['valor_pagado']) for q in qs if q['valor_pagado']) == int(c['total_pagado'])
               and sum(int(q['saldo_cuota']) for q in qs) == int(c['saldo_pendiente'])
               and {q['cliente'] for q in qs} == {c['cliente']})
        seq = sorted(int(q['cuota']) for q in qs)
        okc = okc and seq == list(range(1, len(qs) + 1))
        if not okc:
            malos += 1
            print(f"     detalle inconsistente en crédito #{c['credito']} {c['cliente']}")
    check(malos == 0, f'los 133 créditos cuadran contra su detalle de cuotas ({malos} inconsistentes)')

    print('4. Estados')
    est = defaultdict(int)
    for c in creditos:
        est[c['estado']] += 1
    check(est['CANCELADO'] == ESPERADO['cancelados'], f"CANCELADO: {est['CANCELADO']} == {ESPERADO['cancelados']}")
    check(est['AL DIA'] == ESPERADO['al_dia'], f"AL DIA: {est['AL DIA']} == {ESPERADO['al_dia']}")
    check(est['EN MORA'] == ESPERADO['en_mora'], f"EN MORA: {est['EN MORA']} == {ESPERADO['en_mora']}")
    coher = sum(1 for c in creditos
                if (c['estado'] == 'CANCELADO') == (int(c['saldo_pendiente']) == 0)
                and (c['estado'] == 'EN MORA') == (int(c['cuotas_vencidas']) > 0))
    check(coher == len(creditos), 'estado coherente con saldo y cuotas vencidas en cada crédito')

    print('5. Cuadre por estado de cuota vs fecha de corte')
    mal_estado = 0
    credito_abierto = {num: any(not q['valor_pagado'] for q in qs) for num, qs in por_credito.items()}
    for q in cuotas:
        vencida = q['fecha_vencimiento'] and datetime.date.fromisoformat(q['fecha_vencimiento']) < corte
        if not q['valor_pagado']:
            esperado = 'VENCIDA' if vencida else 'PENDIENTE'
            saldo_ok = int(q['saldo_cuota']) == int(q['valor_letra'])
        elif credito_abierto[q['credito']] and int(q['valor_pagado']) < int(q['valor_letra']):
            esperado = 'ABONO PARCIAL'
            saldo_ok = int(q['saldo_cuota']) == int(q['valor_letra']) - int(q['valor_pagado'])
        else:
            esperado = 'PAGADA'
            saldo_ok = int(q['saldo_cuota']) == 0
        if q['estado'] != esperado or not saldo_ok:
            mal_estado += 1
    check(mal_estado == 0,
          f'estado y saldo de cada cuota consistentes con pago, crédito y corte ({mal_estado} mal)')

    print('6. Cuadre independiente contra el Excel ORIGINAL')
    wb = openpyxl.load_workbook(args.origen, data_only=True)
    ws = wb['LETRAS']
    filas = [r for r in ws.iter_rows(values_only=True) if r[1] is not None]
    check(len(filas) == ESPERADO['cuotas'], f'LETRAS original tiene {len(filas)} filas == {ESPERADO["cuotas"]}')
    orig_pag = round(sum(float(r[5]) for r in filas if r[5] not in (None, '', 0)))
    # saldo recalculado desde el original con la regla de abonos parciales:
    # en créditos con alguna cuota sin pago, un pago menor a la letra deja deuda
    grupos = defaultdict(list)
    for r in filas:
        grupos[r[1]].append(r)
    orig_pend = 0.0
    for qs in grupos.values():
        abierto = any(q[5] in (None, '', 0) for q in qs)
        for q in qs:
            if q[5] in (None, '', 0):
                orig_pend += float(q[4])
            elif abierto and float(q[5]) < float(q[4]):
                orig_pend += float(q[4]) - float(q[5])
    orig_pend = round(orig_pend)
    check(orig_pag == tot_pag, f'Σ pagado ORIGINAL ${orig_pag:,} == Σ pagado consolidado ${tot_pag:,}')
    check(orig_pend == tot_pend, f'Σ saldo ORIGINAL ${orig_pend:,} == Σ saldo consolidado ${tot_pend:,}')

    # pagos de letras en hojas de caja (independiente de consolidar.py)
    total_caja = 0.0
    nombres = {re.sub(r'\s+', ' ', str(r[3]).strip().upper()) for r in filas}
    nombres |= {'NOHORA ELCY MADRIGAL VALENCIA', 'ANYI PAOLA PERAFAN HOYOS'}
    for sn in wb.sheetnames:
        if sn in ('LETRAS', 'NORMA', 'SIMULADOR', 'APORTES', 'Data'):
            continue
        seccion = False
        for r in wb[sn].iter_rows(values_only=True):
            c0 = re.sub(r'\s+', ' ', str(r[0]).strip().upper()) if r[0] is not None else ''
            c2 = re.sub(r'\s+', ' ', str(r[2]).strip().upper()) if len(r) > 2 and r[2] is not None else ''
            c4 = re.sub(r'\s+', ' ', str(r[4]).strip().upper()) if len(r) > 4 and r[4] is not None else ''
            if c0 == 'LETRAS' and 'CUOTA' in c2:
                seccion = True
                continue
            if 'TOTAL INGRESOS' in c2 or 'TOTAL INGRESOS' in c0:
                seccion = False
                continue
            if not seccion or not c4:
                continue
            monto = r[11] if len(r) > 11 else None
            if not isinstance(monto, (int, float)) or monto == 0:
                continue
            es_letra = bool(re.match(r'^(PAGO|PGO)\s+LE[TY]?T?RA', c4)) \
                or c4.startswith('ABONO LETRA') or c4.startswith('SALDO LETRA') \
                or c4.startswith('SALDO DEUDA') or c4 in nombres
            if es_letra:
                total_caja += float(monto)
    check(round(total_caja) == tot_pag,
          f'Σ pagos de letras en 31 hojas de caja ${round(total_caja):,} == consolidado ${tot_pag:,}')

    print('7. Excel consolidado (releído desde disco)')
    wbc = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    check(set(wbc.sheetnames) == {'RESUMEN', 'CLIENTES', 'CREDITOS', 'CUOTAS', 'MORA', 'HALLAZGOS'},
          f'hojas del consolidado: {wbc.sheetnames}')
    check(wbc['CLIENTES'].max_row - 1 == ESPERADO['clientes'], f"hoja CLIENTES: {wbc['CLIENTES'].max_row - 1} filas")
    check(wbc['CREDITOS'].max_row - 1 == ESPERADO['creditos'], f"hoja CREDITOS: {wbc['CREDITOS'].max_row - 1} filas")
    check(wbc['CUOTAS'].max_row - 1 == ESPERADO['cuotas'], f"hoja CUOTAS: {wbc['CUOTAS'].max_row - 1} filas")
    check(wbc['MORA'].max_row - 1 == ESPERADO['en_mora'], f"hoja MORA: {wbc['MORA'].max_row - 1} filas")
    hdr = [c.value for c in wbc['CREDITOS'][1]]
    col_pag, col_saldo = hdr.index('TOTAL_PAGADO') + 1, hdr.index('SALDO_PENDIENTE') + 1
    xl_pag = round(sum(c[0].value or 0 for c in wbc['CREDITOS'].iter_rows(min_row=2, min_col=col_pag, max_col=col_pag)))
    xl_saldo = round(sum(c[0].value or 0 for c in wbc['CREDITOS'].iter_rows(min_row=2, min_col=col_saldo, max_col=col_saldo)))
    check(xl_pag == tot_pag, f'hoja CREDITOS Σ TOTAL_PAGADO ${xl_pag:,} == ${tot_pag:,}')
    check(xl_saldo == tot_pend, f'hoja CREDITOS Σ SALDO_PENDIENTE ${xl_saldo:,} == ${tot_pend:,}')

    print('8. Clientes vs créditos')
    suma_cli = sum(int(c['saldo_pendiente']) for c in clientes)
    check(suma_cli == tot_pend, f'Σ saldo en clientes.csv ${suma_cli:,} == ${tot_pend:,}')
    ids = {c['id'] for c in clientes}
    check(len(ids) == len(clientes), 'IDs de cliente únicos')
    nombres_cli = {c['nombre'] for c in clientes}
    check({c['cliente'] for c in creditos} == nombres_cli, 'mismos clientes en creditos.csv y clientes.csv')

    print()
    if errores:
        print(f'RESULTADO: {len(errores)} VERIFICACIONES FALLARON')
        for e in errores:
            print(f'  - {e}')
        sys.exit(1)
    print('RESULTADO: TODAS LAS VERIFICACIONES PASARON')


if __name__ == '__main__':
    main()
